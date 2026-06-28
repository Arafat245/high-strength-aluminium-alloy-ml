"""Convert data/candidates_aldemo.csv -> data/candidates_aldemo.xlsx
with a 'Candidates' data sheet and a 'ReadMe' column-dictionary sheet.
(Parallel to make_candidates_xlsx.py, but for the ALDEMO Al-Cu-Si-Sc-Zr set.)"""
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
D="/mnt/sdb/arafat/Al_alloy_design/"
df=pd.read_csv(D+"data/candidates_aldemo.csv")
df.insert(0,"#",range(1,len(df)+1))

readme=[
 ("Al","Aluminium - balance, = 100 - sum(other elements)","wt %"),
 ("Cu","Copper - main strengthener; forms theta-prime (Al2Cu) on ageing (2xxx mechanism)","wt %"),
 ("Si","Silicon - in ALDEMO; minor here (improves castability / fluidity). Kept low for wrought","wt %"),
 ("Sc","Scandium - forms Al3Sc / Al3(Sc,Zr) L1-2 dispersoids: grain refinement + thermal stability + strength","wt %"),
 ("Zr","Zirconium - forms Al3Zr / Al3(Sc,Zr); cheap (~$30/kg), routine; anti-recrystallisation","wt %"),
 ("Fe","Iron - impurity, fixed at 0.15 (typical foundry level)","wt %"),
 ("Tsol","Solution(ising) temperature - hold here to dissolve Cu into solid solution, then water-quench","deg C"),
 ("Tage","Ageing temperature - low-T age to precipitate theta-prime (Al2Cu). Al3(Sc,Zr) forms earlier (cast/homogenise)","deg C"),
 ("tage","Ageing time - duration held at Tage","hours"),
 ("UTS","Ultimate tensile strength - PREDICTED, ensemble mean. CONSERVATIVE here (model under-credits Sc/Zr)","MPa"),
 ("UTS s","UTS sigma - std dev across the 12 ensemble models. Honest band is +/-35 MPa (ALDEMO-subset grouped-CV)","MPa"),
 ("YS","Yield strength - PREDICTED, ensemble mean","MPa"),
 ("YS s","YS sigma - std dev across the 12 ensemble models (uncertainty)","MPa"),
 ("EL","Elongation - predicted ductility. LOW reliability (model R2 ~ 0.4); rough guide only","%"),
 ("novelty","Standardised distance to nearest real Al-Cu(-Sc-Zr) alloy. This space is sparse, so 1-2 = genuinely new (validate in Thermo-Calc)","(dimensionless)"),
 ("score","Ranking metric = UTS - 1 x (UTS s). Rewards strength, penalises uncertainty","MPa"),
 ("near_real","The closest real alloy in our data, for context (composition + its measured UTS)","-"),
]
rdf=pd.DataFrame(readme,columns=["Column","Meaning","Units"])

notes=[
 "",
 "WHAT THIS FILE IS:",
 "  5 NEW Al-Cu-Si-Sc-Zr alloys designed to be validated in Thermo-Calc's free ALDEMO demo",
 "  database (elements: Al, Cu, Si, Sc, Zr only - no Zn/Mg, so 7xxx cannot be modelled there).",
 "",
 "HOW TO READ THE HEAT TREATMENT:",
 "  e.g. Tsol 521, Tage 151, tage 35  ->  'Solutionise ~521 C, water-quench, age 151 C for 35 h.'",
 "  NOTE: Al3(Sc,Zr) dispersoids form during casting / homogenisation (~350-450 C), NOT during",
 "  this low-T age - so check their solvus in Thermo-Calc (that is the main reason Sc/Zr are here).",
 "",
 "HOW TO READ A PREDICTION:",
 "  UTS 452, UTS s 13  ->  predicted UTS 452 +/- 13 (ensemble); realistic band 452 +/- 35 MPa.",
 "  Predictions are CONSERVATIVE: real Al-Cu-Zr in our data already reaches 475-510 MPa.",
 "",
 "CAVEATS:",
 "  - Predictions are computational hypotheses; confirm with Thermo-Calc (ALDEMO) + casting/testing.",
 "  - Within +/-35 MPa all 5 are statistically tied in UTS; choose by castability + Sc/Zr cost.",
 "  - Model handles Cu-driven (theta-prime) strength well; it under-credits Al3(Sc,Zr) - their",
 "    benefit (grain refinement, heat resistance) shows up in Thermo-Calc, not in this UTS number.",
 "  - Sc is expensive (~$3000/kg); Zr is cheap. Candidate 5 (lean Sc) is the lowest-cost pick.",
]
rdf2=rdf

with pd.ExcelWriter(D+"data/candidates_aldemo.xlsx",engine="openpyxl") as xl:
    df.to_excel(xl,sheet_name="Candidates",index=False)
    rdf2.to_excel(xl,sheet_name="ReadMe",index=False,startrow=1)
    wb=xl.book
    head=Font(bold=True,color="FFFFFF"); fill=PatternFill("solid",fgColor="305496")
    thin=Border(*[Side(style="thin",color="D9D9D9")]*4)
    rs=wb["ReadMe"]; rs["A1"]="candidates_aldemo.xlsx - column dictionary"; rs["A1"].font=Font(bold=True,size=13)
    for ws in [wb["Candidates"],rs]:
        hdr_row=1 if ws.title=="Candidates" else 2
        for c in range(1,ws.max_column+1):
            cell=ws.cell(hdr_row,c); cell.font=head; cell.fill=fill
            cell.alignment=Alignment(horizontal="center",vertical="center")
        ws.freeze_panes=ws.cell(hdr_row+1,1)
    for c in range(1,wb["Candidates"].max_column+1):
        wb["Candidates"].column_dimensions[get_column_letter(c)].width=10
    wb["Candidates"].column_dimensions[get_column_letter(wb["Candidates"].max_column)].width=26  # near_real
    for c,w in enumerate([12,92,16],1):
        rs.column_dimensions[get_column_letter(c)].width=w
    for r in range(3,3+len(rdf2)):
        rs.cell(r,2).alignment=Alignment(wrap_text=True,vertical="top")
        for c in range(1,4): rs.cell(r,c).border=thin
    start=3+len(rdf2)+1
    for i,line in enumerate(notes):
        cell=rs.cell(start+i,1); cell.value=line
        if line.endswith(":"): cell.font=Font(bold=True)
print("wrote data/candidates_aldemo.xlsx  (sheets: Candidates, ReadMe)")
