"""Convert data/candidates_final.csv -> data/candidates_final.xlsx
with a 'Candidates' data sheet and a 'ReadMe' column-dictionary sheet."""
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
D="/mnt/sdb/arafat/Al_alloy_design/"
df=pd.read_csv(D+"data/candidates_final.csv")
df.insert(0,"#",range(1,len(df)+1))

readme=[
 ("Al","Aluminium - balance, = 100 - sum(other elements)","wt %"),
 ("Zn","Zinc - main strengthening element (forms MgZn2 / eta-prime)","wt %"),
 ("Mg","Magnesium - main strengthening element (pairs with Zn)","wt %"),
 ("Cu","Copper - strengthener; high levels reduce weldability / SCC resistance","wt %"),
 ("Mn","Manganese - cheap dispersoid former (grain-structure control)","wt %"),
 ("Cr","Chromium - cheap dispersoid former (as in 7075)","wt %"),
 ("Fe","Iron - impurity, fixed at 0.15 (typical foundry level)","wt %"),
 ("Si","Silicon - impurity, fixed at 0.10","wt %"),
 ("Tsol","Solution(ising) temperature - hold here to dissolve solute into solid solution, then water-quench (the SHT step)","deg C"),
 ("Tage","Ageing temperature - hold here after quench to precipitate strengthening phases (eta-prime / MgZn2); the T6 age step","deg C"),
 ("tage","Ageing time - duration held at Tage to reach peak strength","hours"),
 ("UTS","Ultimate tensile strength - PREDICTED, mean of the 12-model ensemble","MPa"),
 ("UTS s","UTS sigma - std dev across the 12 ensemble models (model uncertainty). Understates true error; honest band is +/-43 MPa (grouped-CV)","MPa"),
 ("YS","Yield strength - PREDICTED, ensemble mean","MPa"),
 ("YS s","YS sigma - std dev across the 12 ensemble models (uncertainty)","MPa"),
 ("EL","Elongation - predicted ductility. LOW reliability (model R2 ~ 0.3); rough guide only","%"),
 ("novelty","Standardised distance to nearest real 7xxx alloy. <=1.32 = in-distribution / reliable","(dimensionless)"),
 ("score","Ranking metric = UTS - 1 x (UTS s). Rewards strength, penalises uncertainty. File sorted by this","MPa"),
]
rdf=pd.DataFrame(readme,columns=["Column","Meaning","Units"])

notes=[
 "",
 "HOW TO READ THE HEAT TREATMENT (T6 temper):",
 "  e.g. Tsol 488, Tage 104, tage 24  ->  'Solutionise at 488 C, water-quench, then age at 104 C for 24 h.'",
 "",
 "HOW TO READ A PREDICTION:",
 "  UTS 635, UTS s 13.6  ->  predicted UTS 635 +/- 14 MPa (ensemble); realistic band 635 +/- 43 MPa.",
 "",
 "CAVEATS:",
 "  - Predictions are computational hypotheses (+/-43 MPa); confirm with Thermo-Calc + casting/testing.",
 "  - Within +/-43 MPa, candidates 1-3 are statistically tied in UTS; use YS and castability to choose.",
 "  - The 's' suffix = sigma (std dev across 12 bootstrap models), an uncertainty estimate, not a property.",
 "  - Model does not see grain size or multi-stage ageing (RRA/T74).",
]

with pd.ExcelWriter(D+"data/candidates_final.xlsx",engine="openpyxl") as xl:
    df.to_excel(xl,sheet_name="Candidates",index=False)
    rdf.to_excel(xl,sheet_name="ReadMe",index=False,startrow=1)
    wb=xl.book
    head=Font(bold=True,color="FFFFFF"); fill=PatternFill("solid",fgColor="305496")
    thin=Border(*[Side(style="thin",color="D9D9D9")]*4)
    # ReadMe title
    rs=wb["ReadMe"]; rs["A1"]="candidates_final.xlsx - column dictionary"; rs["A1"].font=Font(bold=True,size=13)
    for ws,widths in [(wb["Candidates"],None),(rs,[12,86,16])]:
        hdr_row=1 if ws.title=="Candidates" else 2
        for c in range(1,ws.max_column+1):
            cell=ws.cell(hdr_row,c); cell.font=head; cell.fill=fill
            cell.alignment=Alignment(horizontal="center",vertical="center")
        ws.freeze_panes=ws.cell(hdr_row+1,1)
    # widths
    for c in range(1,wb["Candidates"].max_column+1):
        wb["Candidates"].column_dimensions[get_column_letter(c)].width=10
    for c,w in enumerate([12,86,16],1):
        rs.column_dimensions[get_column_letter(c)].width=w
    for r in range(3,3+len(rdf)):
        rs.cell(r,2).alignment=Alignment(wrap_text=True,vertical="top")
        for c in range(1,4): rs.cell(r,c).border=thin
    # notes below table
    start=3+len(rdf)+1
    for i,line in enumerate(notes):
        cell=rs.cell(start+i,1); cell.value=line
        if line.endswith(":"): cell.font=Font(bold=True)
print("wrote data/candidates_final.xlsx  (sheets: Candidates, ReadMe)")
